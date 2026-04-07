"""
Xử lý tài liệu: extract text → chunk → embed → lưu vào DB.
"""
import io
import uuid
from pathlib import Path

from sqlalchemy.ext.asyncio import AsyncSession

from app.models.knowledge import DocChunk, KnowledgeDoc
from app.services.embedding import embed_batch

CHUNK_SIZE = 512    # tokens (tương đương ~400 words)
CHUNK_OVERLAP = 64  # tokens


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def extract_text_from_pdf(data: bytes) -> str:
    from PyPDF2 import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_text_from_xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(c) for c in row if c is not None)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)


def extract_text(data: bytes, file_type: str) -> str:
    if file_type == "pdf":
        return extract_text_from_pdf(data)
    elif file_type == "docx":
        return extract_text_from_docx(data)
    elif file_type in ("xlsx", "xls"):
        return extract_text_from_xlsx(data)
    elif file_type in ("md", "txt"):
        return data.decode("utf-8", errors="replace")
    raise ValueError(f"Loại file không hỗ trợ: {file_type}")


# ---------------------------------------------------------------------------
# Chunking (character-based approximation, ~4 chars/token)
# ---------------------------------------------------------------------------

CHARS_PER_TOKEN = 4

def chunk_text(text: str) -> list[str]:
    chunk_chars = CHUNK_SIZE * CHARS_PER_TOKEN
    overlap_chars = CHUNK_OVERLAP * CHARS_PER_TOKEN
    step = chunk_chars - overlap_chars

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_chars
        chunks.append(text[start:end].strip())
        start += step

    return [c for c in chunks if len(c) > 50]  # lọc chunk quá ngắn


# ---------------------------------------------------------------------------
# Full pipeline: raw bytes → DB
# ---------------------------------------------------------------------------

async def process_and_store(
    db: AsyncSession,
    doc: KnowledgeDoc,
    file_data: bytes,
    file_type: str,
) -> int:
    """
    Extract text từ file, chunk, embed và lưu doc_chunks.
    Trả về số lượng chunks đã lưu.
    """
    raw_text = extract_text(file_data, file_type)
    doc.content_raw = raw_text[:50_000]  # giới hạn lưu raw text

    chunks = chunk_text(raw_text)
    embeddings = await embed_batch(chunks)

    db_chunks = []
    for i, (chunk_text_val, embedding) in enumerate(zip(chunks, embeddings)):
        db_chunks.append(
            DocChunk(
                id=uuid.uuid4(),
                doc_id=doc.id,
                chunk_index=i,
                chunk_text=chunk_text_val,
                embedding=embedding,
                token_count=len(chunk_text_val) // CHARS_PER_TOKEN,
            )
        )

    db.add_all(db_chunks)
    return len(db_chunks)


# ---------------------------------------------------------------------------
# RAG search
# ---------------------------------------------------------------------------

async def rag_search(
    db: AsyncSession,
    query: str,
    user_roles: list[str],
    client_id: uuid.UUID | None = None,
    limit: int = 5,
    score_threshold: float = 0.7,
) -> list[dict]:
    """
    Embed query → cosine similarity trong pgvector → trả top chunks.
    """
    from sqlalchemy import text

    from app.services.embedding import embed_text

    query_vec = await embed_text(query)
    vec_str = "[" + ",".join(str(v) for v in query_vec) + "]"

    sql = text("""
        SELECT
            dc.chunk_text,
            kd.title,
            kd.category,
            kd.source_url,
            1 - (dc.embedding <=> CAST(:vec AS vector)) AS score
        FROM doc_chunks dc
        JOIN knowledge_docs kd ON dc.doc_id = kd.id
        WHERE kd.is_active = TRUE
          AND kd.access_level = ANY(:roles)
          AND (CAST(:client_id AS uuid) IS NULL OR kd.client_id = CAST(:client_id AS uuid) OR kd.client_id IS NULL)
        ORDER BY dc.embedding <=> CAST(:vec AS vector)
        LIMIT :limit
    """)

    result = await db.execute(
        sql,
        {
            "vec": vec_str,
            "roles": user_roles,
            "client_id": str(client_id) if client_id else None,  # cast trong SQL
            "limit": limit,
        },
    )
    rows = result.fetchall()

    return [
        {
            "chunk_text": r.chunk_text,
            "title": r.title,
            "category": r.category,
            "source_url": r.source_url,
            "score": float(r.score),
        }
        for r in rows
        if float(r.score) >= score_threshold
    ]
